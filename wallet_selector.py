from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime

class WalletSelector:
    def __init__(self, excel_path):
        self.excel_path = Path(excel_path)
        self.wb = load_workbook(self.excel_path)
        self.ws = self.wb.active
        self.headers = {}
        for cell in self.ws[1]:
            self.headers[str(cell.value).strip()] = cell.column

    def _cell(self, row, header):
        return self.ws.cell(row=row, column=self.headers[header])

    def _max_row(self):
        return self.ws.max_row

    def get_random_incomplete_wallet(self):
        rows = []
        for r in range(2, self._max_row() + 1):
            name = str(self._cell(r, "name").value or "").strip()
            if not name or "last update" in name.lower():
                continue
            done = int(self._cell(r, "Done").value or 0)
            if done == 1:
                continue
            holesky = int(self._cell(r, "HoleskyTransaction").value or 0)
            babylon = int(self._cell(r, "BabylonTransaction").value or 0)
            xion = int(self._cell(r, "XionTransaction").value or 0)
            sei = int(self._cell(r, "SeiTransaction").value or 0)
            rows.append({
                "row": r,
                "name": name,
                "HoleskyTransaction": holesky,
                "BabylonTransaction": babylon,
                "XionTransaction": xion,
                "SeiTransaction": sei
            })
        if not rows:
            return None
        import random
        return random.choice(rows)

    def update_wallet(self, wallet_name, holesky, babylon, xion, sei):
        for r in range(2, self._max_row() + 1):
            name = str(self._cell(r, "name").value or "").strip()
            if name == str(wallet_name).strip():
                self._cell(r, "HoleskyTransaction").value = int(holesky)
                self._cell(r, "BabylonTransaction").value = int(babylon)
                self._cell(r, "XionTransaction").value = int(xion)
                self._cell(r, "SeiTransaction").value = int(sei)
                return

    def mark_wallet_done(self, wallet_name):
        for r in range(2, self._max_row() + 1):
            name = str(self._cell(r, "name").value or "").strip()
            if name == str(wallet_name).strip():
                self._cell(r, "Done").value = 1
                return

    def save(self):
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.ws["A1"] = self.ws["A1"].value
        self.wb.save(self.excel_path)
        try:
            from openpyxl import load_workbook
            wb = load_workbook(self.excel_path)
            ws = wb.active
            ws.append([f"Last update: {timestamp}"])
            wb.save(self.excel_path)
        except Exception:
            pass