from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

REQUIRED_HEADERS = (
    "name",
    "HoleskyTransaction",
    "BabylonTransaction",
    "XionTransaction",
    "SeiTransaction",
    "Done",
)

# Fixed cell used to stamp the last-saved time. Kept well clear of the
# tracked columns (A-F) so it never collides with real wallet data.
_LAST_UPDATE_CELL = "H1"


class WalletSelector:
    def __init__(self, excel_path):
        self.excel_path = Path(excel_path)
        if not self.excel_path.exists():
            raise FileNotFoundError(
                f"Wallet spreadsheet not found: {self.excel_path}. "
                "Create it with the columns described in the README."
            )
        self.wb = load_workbook(self.excel_path)
        self.ws = self.wb.active
        self.headers = {}
        for cell in self.ws[1]:
            if cell.value:
                self.headers[str(cell.value).strip()] = cell.column
        missing = [h for h in REQUIRED_HEADERS if h not in self.headers]
        if missing:
            raise ValueError(
                f"{self.excel_path} is missing required column(s): {missing}"
            )

    def _cell(self, row, header):
        return self.ws.cell(row=row, column=self.headers[header])

    def _max_row(self):
        return self.ws.max_row

    def get_random_incomplete_wallet(self):
        rows = []
        for r in range(2, self._max_row() + 1):
            name = str(self._cell(r, "name").value or "").strip()
            if not name or "last update" in name.lower():
                continue
            done = int(self._cell(r, "Done").value or 0)
            if done == 1:
                continue
            rows.append({
                "row": r,
                "name": name,
                "HoleskyTransaction": int(self._cell(r, "HoleskyTransaction").value or 0),
                "BabylonTransaction": int(self._cell(r, "BabylonTransaction").value or 0),
                "XionTransaction": int(self._cell(r, "XionTransaction").value or 0),
                "SeiTransaction": int(self._cell(r, "SeiTransaction").value or 0),
            })
        if not rows:
            return None
        import random
        return random.choice(rows)

    def update_wallet(self, wallet_name, holesky, babylon, xion, sei):
        for r in range(2, self._max_row() + 1):
            name = str(self._cell(r, "name").value or "").strip()
            if name == str(wallet_name).strip():
                self._cell(r, "HoleskyTransaction").value = int(holesky)
                self._cell(r, "BabylonTransaction").value = int(babylon)
                self._cell(r, "XionTransaction").value = int(xion)
                self._cell(r, "SeiTransaction").value = int(sei)
                return
        raise ValueError(f"Wallet '{wallet_name}' not found in spreadsheet")

    def mark_wallet_done(self, wallet_name):
        for r in range(2, self._max_row() + 1):
            name = str(self._cell(r, "name").value or "").strip()
            if name == str(wallet_name).strip():
                self._cell(r, "Done").value = 1
                return
        raise ValueError(f"Wallet '{wallet_name}' not found in spreadsheet")

    def save(self):
        """Persist the workbook.

        Fix: the previous implementation reopened the file and *appended* a
        fresh "Last update: ..." row on every single save (i.e. after every
        wallet processed). Left running for a while, this silently added
        hundreds of junk rows and made the spreadsheet grow without bound.
        The timestamp is now written to one fixed, reserved cell instead.
        """
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.ws[_LAST_UPDATE_CELL] = f"Last update: {timestamp}"
        self.wb.save(self.excel_path)
